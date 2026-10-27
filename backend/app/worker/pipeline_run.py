"""Pure, testable helpers for the pipeline worker.

Kept free of DB/S3/Celery so they can be unit-tested without infrastructure.
"""
from __future__ import annotations

import csv
import gzip
import io
import os
from dataclasses import dataclass

from app.models.enums import ResultKind


def count_fastq_reads(path: str, stop_at: int | None = None) -> int:
    """Count reads (4 lines each) in a (optionally gzipped) FASTQ.

    If ``stop_at`` is given, stop once that many reads are reached and return it — a cheap way to
    answer "are there at least N reads?" without decompressing a whole 2 GB file when it's fine.
    """
    limit = None if stop_at is None else stop_at * 4

    def _count_lines(fh) -> int:
        n = 0
        for _ in fh:
            n += 1
            if limit is not None and n >= limit:
                break
        return n

    try:
        with gzip.open(path, "rt") as fh:
            lines = _count_lines(fh)
    except gzip.BadGzipFile:  # tolerate a plain (uncompressed) FASTQ
        with open(path, "rt") as fh:
            lines = _count_lines(fh)
    return lines // 4

# Map published pipeline output filenames (suffixes) to a ResultKind.
_RESULT_SUFFIXES = {
    "_consensus_genotypes.txt": ResultKind.consensus,  # must win over "_genotypes.txt"
    "_reference_alleles.txt": ResultKind.reference_alleles,
    "_genotypes.txt": ResultKind.genotypes,
    "_positions.txt": ResultKind.positions,
    "_frequency_of_sequences_by_marker.txt": ResultKind.frequency,
    "_reads_summary.csv": ResultKind.reads_summary,
    "_consensus_report.html": ResultKind.consensus_report,  # must win over "_report.html"
    "_report.html": ResultKind.html_report,
}


def build_nextflow_cmd(
    *, pipeline_dir: str, input_tsv: str, run_dir: str, profile: str,
    min_identity: float, min_overlap: int, expected_read_number: int | None = None,
    parameters_file_path: str | None = None,
) -> list[str]:
    """Argument vector for `nextflow run`, meant to run with cwd=run_dir.

    We deliberately do NOT pass --outdir: main.nf defaults params.outdir to the kit_id
    and nests results/ + reports/ under it, so outputs land at {run_dir}/{kit_id}/...
    (see collect_results). Runs headless: no ANSI, log inside the run dir.
    """
    cmd = [
        "nextflow", "-log", os.path.join(run_dir, ".nextflow.log"),
        "run", os.path.join(pipeline_dir, "main.nf"),
        "-profile", profile,
        "-ansi-log", "false",
        "--input", input_tsv,
        "--min_identity", str(min_identity),
        "--min_overlap", str(min_overlap),
    ]
    if expected_read_number:  # drives the reference line in the HTML report
        cmd += ["--expected_read_number", str(expected_read_number)]
    if parameters_file_path:  # per-run pipeline thresholds (else scripts use the baked default)
        cmd += ["--parameters_file_path", parameters_file_path]
    return cmd


@dataclass
class CollectedResult:
    kind: ResultKind
    path: str
    filename: str


def archive_logs(scratch: str, kit_id: str, dest_path: str) -> str | None:
    """Tar up everything that explains a run: the per-locus logs the pipeline publishes under
    reports/logs/, Nextflow's own .nextflow.log, and our captured run.log.

    The scratch directory is deleted after a successful run, so without this a run that *looked*
    fine (a locus silently emptied by a collapsed threshold, say) leaves nothing to inspect.
    Returns the archive path, or None when there was nothing to collect.
    """
    import tarfile

    candidates = [
        (os.path.join(scratch, kit_id, "reports", "logs"), "logs"),
        (os.path.join(scratch, ".nextflow.log"), ".nextflow.log"),
        (os.path.join(scratch, "run.log"), "run.log"),
    ]
    present = [(path, arcname) for path, arcname in candidates if os.path.exists(path)]
    if not present:
        return None
    with tarfile.open(dest_path, "w:gz") as tar:
        for path, arcname in present:
            tar.add(path, arcname=arcname)
    return dest_path


def collect_results(outdir: str, kit_id: str) -> list[CollectedResult]:
    """Find published outputs under {outdir}/{kit_id}/{results,reports}.

    Mirrors the pipeline's publishDir layout (results_dir / reports_dir).
    """
    found: list[CollectedResult] = []
    for sub in ("results", "reports"):
        d = os.path.join(outdir, kit_id, sub)
        if not os.path.isdir(d):
            continue
        for name in sorted(os.listdir(d)):
            # Longest suffix first so "_consensus_genotypes.txt" beats "_genotypes.txt".
            for suffix, kind in sorted(_RESULT_SUFFIXES.items(), key=lambda kv: -len(kv[0])):
                if name.endswith(suffix):
                    found.append(CollectedResult(kind, os.path.join(d, name), name))
                    break
    return found


def find_result(results: list[CollectedResult], kind: ResultKind) -> str | None:
    for r in results:
        if r.kind == kind:
            return r.path
    return None


def samples_text_to_rows(text: str) -> list[dict]:
    """Parse pasted sample text into TPositionId/SPositionBC(/control_type) rows.

    Accepts CSV/TSV/whitespace with two or three fields per line: position (e.g. A1), sample name,
    and an optional control_type (empty for ordinary samples). A header line containing 'position'
    (any case) is skipped.
    """
    rows: list[dict] = []
    for raw in text.strip().splitlines():
        line = raw.strip()
        if not line:
            continue
        parts = [p.strip() for p in line.replace("\t", ",").split(",")]
        # drop trailing empties but keep an interior empty control_type field
        while parts and parts[-1] == "":
            parts.pop()
        if len([p for p in parts if p]) < 2:
            parts = line.split()
        if len(parts) < 2:
            raise ValueError(f"cannot parse sample line: {raw!r} (need position and name)")
        pos, name = parts[0], parts[1]
        control_type = parts[2].strip() if len(parts) >= 3 else ""
        if pos.lower() == "tpositionid" or pos.lower() == "position":
            continue
        rows.append({"TPositionId": pos, "SPositionBC": name, "control_type": control_type})
    if not rows:
        raise ValueError("no sample rows parsed from pasted text")
    return rows


FULL_PLATE_WELLS = 96  # a batch is one physical 96-well plate — every well must be filled


def count_filled_wells(rows: list[dict]) -> int:
    """Distinct wells that carry a sample name, from parsed TPositionId/SPositionBC rows."""
    seen: set[str] = set()
    for r in rows:
        pos = (r.get("TPositionId") or "").strip().upper()
        name = (r.get("SPositionBC") or "").strip()
        if pos and name:
            seen.add(pos)
    return len(seen)


def read_sample_xlsx(path: str) -> list[dict]:
    """Read a TPositionId/SPositionBC(/control_type) .xlsx into rows (header skipped)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows: list[dict] = []
        for cells in ws.iter_rows(values_only=True):
            if not cells:
                continue
            pos = str(cells[0]).strip() if cells[0] is not None else ""
            name = str(cells[1]).strip() if len(cells) > 1 and cells[1] is not None else ""
            ct = str(cells[2]).strip() if len(cells) > 2 and cells[2] is not None else ""
            if pos.lower() in ("tpositionid", "position"):  # header row
                continue
            if not pos and not name:
                continue
            rows.append({"TPositionId": pos, "SPositionBC": name, "control_type": ct})
        return rows
    finally:
        wb.close()


def write_sample_xlsx(rows: list[dict], dest_path: str) -> None:
    """Write TPositionId/SPositionBC/control_type rows to an .xlsx make_ngsfilter.py reads."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["TPositionId", "SPositionBC", "control_type"])
    for r in rows:
        ws.append([r["TPositionId"], r["SPositionBC"], r.get("control_type", "")])
    wb.save(dest_path)


def sample_rows_to_csv_preview(rows: list[dict]) -> str:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=["TPositionId", "SPositionBC", "control_type"],
                       extrasaction="ignore")
    w.writeheader()
    w.writerows(rows)
    return buf.getvalue()
